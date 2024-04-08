
import os
import time
import json as ujson
import shutil
import requests

from io import BytesIO
from docx.shared import Cm
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from concurrent.futures import ThreadPoolExecutor, wait


stat = {
}

sess = requests.Session()

def edit_margin(section, instruct):
    section.left_margin = Cm(instruct["margins"]["left"])
    section.right_margin = Cm(instruct["margins"]["right"])
    section.top_margin = Cm(instruct["margins"]["top"])
    section.bottom_margin = Cm(instruct["margins"]["bottom"])

def add_img(cell, img, width, height):
    paragraph = cell.paragraphs[0]
    # Add image
    run = paragraph.add_run()
    run.add_picture(img, width=width, height=height)

def wrap_img(binary):
    return BytesIO(binary)

def download_images(document, codes, images, img_width, img_height):

    img_row = document.sections[0].page_width // img_width
    table = document.add_table(rows=(images + img_row - 1) // img_row, cols=img_row)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    dif_imgs = len(codes)

    counter = [0]

    def mapper(i, index):
        try:
            content = sess.get(codes[i][0]).content

            for _ in range(codes[i][1]):
                add_img(table.rows[index // img_row].cells[index % img_row], wrap_img(content), img_width, img_height)
                index += 1
        except:
            pass
        counter[0] += 1

    with ThreadPoolExecutor(max_workers=10) as exe:
        
        f = []
        index = 0

        for i in range(dif_imgs):
            f.append(exe.submit(mapper, i, index))
            index += codes[i][1]

    wait(f)

def parse_ids(raw: str):

    return None, None

def read_instruct(fp):
    if fp == None:
        return fp
    
    data = ujson.load(fp)

    return {
        "width": data["width"],
        "height": data["height"],
        "margins": {
            "left": data["margins"]["left"],
            "right": data["margins"]["right"],
            "top": data["margins"]["top"],
            "bottom": data["margins"]["bottom"]
        } if data.get("margins") != None else None
    }

def proc_file(data, instruct, output):
  
    stat['name'] = output

    codes, images = parse_ids(data)

    img_w = Cm(instruct["width"])
    img_h = Cm(instruct["height"])

    document = Document()
    section = document.sections[-1]

    edit_margin(section, instruct)

    download_images(document, codes, images, img_w, img_h)

    document.save(f"{output}.docx")
    

def main(deck, instruct, output):
    try:
        proc_file(deck, instruct, output)
    except:
        return False
    return True